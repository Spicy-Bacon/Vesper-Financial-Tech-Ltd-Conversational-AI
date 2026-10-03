"""Import only Questions, Answer_Options and Approved Explanations.

Run from the repository root with python -m backend.scripts.import_finance_workbook.
Requires openpyxl 3.1.5 at import time only. Runtime JSON readers do not need it.
Catalog approval defaults to false and is never inferred from the filename.
"""

import argparse
import hashlib
from io import BytesIO
import os
from pathlib import Path
import shutil
import sys
import tempfile
from zipfile import BadZipFile

from backend.app.catalog.release import WORKBOOK_NAME, Provenance, deterministic_json, validate_documents

# Inspected Judge_Ready headers. Select by name, never position. Ignore scoring
# columns and never select any sheet outside this allowlist.
FIELDS = {
    "Questions": ("question_id", "display_order", "category", "question_text", "help_text",
                  "is_safety_question", "catalog_version"),
    "Answer_Options": ("option_id", "question_id", "display_order", "option_label", "confirmation_text",
                       "is_uncertain"),
    "Explanations": ("content_id", "question_id", "explanation_text", "review_status", "catalog_version"),
}


def _text(value, location):
    if not isinstance(value, str) or not value.strip():
        raise ValueError(f"{location}: required nonblank text")
    if value != value.strip():
        raise ValueError(f"{location}: surrounding whitespace requires source correction")
    return value


def _order(value, location):
    if type(value) is not int or value < 1:
        raise ValueError(f"{location}: display_order must be a positive integer")
    return value


def _rows(workbook, sheet_name):
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"missing required sheet: {sheet_name}")
    sheet = workbook[sheet_name]
    if sheet.max_row is None or sheet.max_column is None:
        sheet.calculate_dimension(force=True)
    if sheet.max_row > 10000 or sheet.max_column > 100:
        raise ValueError(f"{sheet_name}: exceeds bounded import size")
    headers = [cell.value for cell in sheet[1]]
    present = [h for h in headers if h is not None]
    if len(present) != len(set(present)):
        raise ValueError(f"{sheet_name}: duplicate headers")
    missing = set(FIELDS[sheet_name]) - set(present)
    if missing:
        raise ValueError(f"{sheet_name}: missing headers {', '.join(sorted(missing))}")
    columns = {field: headers.index(field) for field in FIELDS[sheet_name]}
    for row_number, row in enumerate(sheet.iter_rows(min_row=2), start=2):
        selected = {field: row[index] for field, index in columns.items()}
        if all(cell.value is None for cell in selected.values()):
            continue
        if any(cell.data_type in {"f", "e"} for cell in selected.values()):
            raise ValueError(f"{sheet_name} row {row_number}: formula/error in runtime source field")
        yield row_number, {field: cell.value for field, cell in selected.items()}


def parse_workbook(raw: bytes, *, finance_approved=False, approval_basis=None):
    """Build and validate both envelopes in memory; no output files written."""
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise ValueError("import requires openpyxl==3.1.5; install it in the importer environment") from None
    provenance = Provenance(workbook_filename=WORKBOOK_NAME,
                            workbook_sha256=hashlib.sha256(raw).hexdigest()).model_dump()
    workbook = load_workbook(BytesIO(raw), read_only=True, data_only=False, keep_links=False)
    try:
        questions, orders, version = {}, set(), None
        for row, data in _rows(workbook, "Questions"):
            loc = f"Questions row {row}"
            qid = _text(data["question_id"], loc)
            order = _order(data["display_order"], loc)
            if qid in questions or order in orders:
                raise ValueError(f"{loc}: duplicate question ID or display order")
            orders.add(order)
            current_version = _text(data["catalog_version"], loc)
            if version is not None and version != current_version:
                raise ValueError(f"{loc}: inconsistent catalog versions")
            version = current_version
            if type(data["is_safety_question"]) is not bool:
                raise ValueError(f"{loc}: invalid safety flag; expected an Excel boolean")
            for field in ("category", "question_text", "help_text"):
                _text(data[field], f"{loc} {field}")
            questions[qid] = data
        if "Q7" not in questions or len(questions) < 2:
            raise ValueError("Questions must contain answer questions and a separate Q7")

        options = {qid: [] for qid in questions}
        option_ids, option_orders = set(), set()
        for row, data in _rows(workbook, "Answer_Options"):
            loc = f"Answer_Options row {row}"
            oid = _text(data["option_id"], loc)
            qid = _text(data["question_id"], loc)
            order = _order(data["display_order"], loc)
            if oid in option_ids or (qid, order) in option_orders:
                raise ValueError(f"{loc}: duplicate option ID or display order")
            if qid not in questions:
                raise ValueError(f"{loc}: unknown question reference")
            option_ids.add(oid)
            option_orders.add((qid, order))
            _text(data["option_label"], loc)
            _text(data["confirmation_text"], loc)
            if type(data["is_uncertain"]) is not bool:
                raise ValueError(f"{loc}: invalid is_uncertain flag; expected an Excel boolean")
            options[qid].append(data)

        projected = []
        for qid, data in sorted(questions.items(), key=lambda item: item[1]["display_order"]):
            choices = sorted(options[qid], key=lambda o: o["display_order"])
            if not choices:
                raise ValueError(f"{qid}: missing required options")
            projected.append({
                "id": qid, "label": data["category"], "prompt": data["question_text"],
                "clarification": data["help_text"] + "\n" + "\n".join(o["option_label"] for o in choices),
                "requiresExplicitConfirmation": data["is_safety_question"],
                "options": [{"id": o["option_id"], "answer": o["confirmation_text"],
                             "label": o["option_label"], "is_unsure": o["is_uncertain"]}
                            for o in choices],
            })

        records, content_ids = [], set()
        for row, data in _rows(workbook, "Explanations"):
            loc = f"Explanations row {row}"
            cid = _text(data["content_id"], loc)
            qid = _text(data["question_id"], loc)
            if cid in content_ids:
                raise ValueError(f"{loc}: duplicate explanation ID")
            content_ids.add(cid)
            if qid not in questions:
                raise ValueError(f"{loc}: unknown question reference")
            if data["catalog_version"] != version:
                raise ValueError(f"{loc}: inconsistent catalog version")
            _text(data["explanation_text"], loc)
            status = _text(data["review_status"], loc)
            if status not in {"Approved", "Draft", "Unapproved", "Retired"}:
                raise ValueError(f"{loc}: unknown review_status")
            if status != "Approved":
                continue
            records.append({"content_id": cid, "question_id": qid, "text": data["explanation_text"],
                            "source_sheet": "Explanations", "source_row": row,
                            "catalog_version": version, "approval_status": "approved",
                            "retrievable": qid != "Q7"})
        return validate_documents({
            "schema_version": "finance-catalog-release-v1", "provenance": provenance,
            "approval": {"approved": finance_approved, "basis": approval_basis},
            "catalog": {"version": version, "financeApproved": finance_approved, "demo": False,
                        "questions": [q for q in projected if q["id"] != "Q7"]},
            "question_help": {qid: data["help_text"] for qid, data in questions.items()},
            "review": next(q for q in projected if q["id"] == "Q7"),
        }, {"schema_version": "finance-explanations-release-v1", "provenance": provenance,
            "catalog_version": version, "records": records})
    finally:
        workbook.close()


def import_workbook(source, *, output_root="data/catalog", source_root="data/source",
                    finance_approved=False, approval_basis=None):
    source = Path(source)
    if source.name != WORKBOOK_NAME:
        raise ValueError(f"exact source required: {WORKBOOK_NAME}")
    if not source.is_file():
        raise FileNotFoundError(f"required source is unavailable: {WORKBOOK_NAME}")
    raw = source.read_bytes()
    catalog, explanations = parse_workbook(raw, finance_approved=finance_approved,
                                          approval_basis=approval_basis)
    payloads = {"catalog.json": deterministic_json(catalog),
                "explanations.json": deterministic_json(explanations)}
    output_root = Path(output_root)
    target = output_root / catalog.catalog.version
    if target.exists() and any(not (target / name).is_file() or
                               (target / name).read_bytes() != body for name, body in payloads.items()):
        raise ValueError("different release already exists under this catalog version; refusing overwrite")
    archive = Path(source_root) / catalog.provenance.workbook_sha256 / WORKBOOK_NAME
    if archive.exists() and archive.read_bytes() != raw:
        raise ValueError("source archive differs from its SHA-256; refusing overwrite")
    archive.parent.mkdir(parents=True, exist_ok=True)
    try:
        with archive.open("xb") as stream:
            stream.write(raw)
    except FileExistsError:
        if archive.read_bytes() != raw:
            raise ValueError("source archive changed during import") from None
    if not target.exists():
        output_root.mkdir(parents=True, exist_ok=True)
        staging = Path(tempfile.mkdtemp(prefix=".import-", dir=output_root))
        try:
            for name, body in payloads.items():
                (staging / name).write_bytes(body)
            # Both documents appear together. A non-empty concurrent release
            # cannot be overwritten by this directory rename.
            os.rename(staging, target)
        finally:
            if staging.exists():
                shutil.rmtree(staging)
    return target


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--output-root", type=Path, default=Path("data/catalog"))
    parser.add_argument("--source-root", type=Path, default=Path("data/source"))
    parser.add_argument("--finance-approved", action="store_true")
    parser.add_argument("--approval-basis", help="explicit trusted approval record, required with --finance-approved")
    args = parser.parse_args(argv)
    try:
        target = import_workbook(args.workbook, output_root=args.output_root, source_root=args.source_root,
                                 finance_approved=args.finance_approved, approval_basis=args.approval_basis)
    except (ValueError, OSError, BadZipFile) as error:
        print(f"Import blocked: {error}", file=sys.stderr)
        return 2
    print(f"Imported {target}; financeApproved={args.finance_approved}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
