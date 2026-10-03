"""Real source import tests and mocked transport integration, never live quality claims.

Edited workbooks below are temporary test fixtures, not candidate releases.
"""

from io import BytesIO
import json
from pathlib import Path

import httpx
from openpyxl import load_workbook
import pytest
from fastapi.testclient import TestClient

from backend.app.adapters.omlx import OmlxAdapter
from backend.app.catalog import JsonCatalogProvider
from backend.app.catalog.release import WORKBOOK_NAME, deterministic_json, load_release
from backend.app.main import create_app
from backend.app.services.retrieval import CatalogRetriever
from backend.scripts.import_finance_workbook import FIELDS, import_workbook, parse_workbook
from backend.tests.conftest import Journey

ROOT = Path(__file__).resolve().parents[2]
RELEASE = ROOT / "data/catalog/catalog-v2"


@pytest.fixture
def real_source():
    provider = JsonCatalogProvider(RELEASE)
    path = ROOT / "data/source" / provider.provenance.workbook_sha256 / WORKBOOK_NAME
    assert path.is_file(), "archived authoritative source must exist"
    return path


def edited_bytes(path, edit):
    workbook = load_workbook(path, data_only=False)
    edit(workbook)
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def test_real_import_deterministic_order_wording_and_review(real_source, tmp_path):
    release, explanations = load_release(RELEASE)
    kwargs = dict(finance_approved=release.approval.approved, approval_basis=release.approval.basis)
    first = import_workbook(real_source, output_root=tmp_path / "catalog", source_root=tmp_path / "source", **kwargs)
    before = {p.name: p.read_bytes() for p in first.iterdir()}
    assert import_workbook(real_source, output_root=tmp_path / "catalog", source_root=tmp_path / "source", **kwargs) == first
    assert {p.name: p.read_bytes() for p in first.iterdir()} == before
    assert before == {name: (RELEASE / name).read_bytes() for name in before}
    archived = tmp_path / "source" / release.provenance.workbook_sha256 / WORKBOOK_NAME
    assert archived.read_bytes() == real_source.read_bytes()
    provider = JsonCatalogProvider(first)
    catalog = provider.load()
    assert [q.id for q in catalog.questions] == [f"Q{i}" for i in range(1, 7)]
    assert sum(len(q.options) for q in catalog.questions) == 32
    source = load_workbook(real_source, read_only=False, data_only=False)
    try:
        question_rows = {row[0]: row for row in list(source["Questions"].values)[1:]}
        option_rows = {row[0]: row for row in list(source["Answer_Options"].values)[1:]}
        for question in [*catalog.questions, provider.review]:
            original = question_rows[question.id]
            assert question.prompt == original[3]
            assert provider.question_help[question.id] == original[4]
            assert question.requiresExplicitConfirmation is original[7]
            expected = sorted((r for r in option_rows.values() if r[1] == question.id), key=lambda r: r[2])
            assert [o.id for o in question.options] == [r[0] for r in expected]
            assert [o.answer for o in question.options] == [r[4] for r in expected]
            assert question.clarification == original[4] + "\n" + "\n".join(r[3] for r in expected)
        rows = list(source["Explanations"].values)
        assert len(explanations.records) == 17
        assert sum(r.retrievable for r in explanations.records) == 10
        for record in explanations.records:
            original = rows[record.source_row - 1]
            assert (record.content_id, record.question_id, record.text) == (original[0], original[1], original[3])
    finally:
        source.close()


def test_only_allowlisted_fields_affect_runtime_content(real_source):
    def contaminate(workbook):
        for sheet in workbook:
            if sheet.title not in FIELDS:
                sheet["A1"] = "FORBIDDEN_EXPECTATIONS_SENTINEL"
        # Even embedded scoring columns on an allowed sheet are not selected.
        for row in workbook["Answer_Options"].iter_rows(min_row=2):
            for cell in row[6:]:
                cell.value = "FORBIDDEN_SCORING_SENTINEL"
    baseline = parse_workbook(real_source.read_bytes())
    modified = parse_workbook(edited_bytes(real_source, contaminate))
    for a, b in zip(baseline, modified):
        assert a.model_dump(exclude={"provenance"}) == b.model_dump(exclude={"provenance"})
        text = deterministic_json(b).decode()
        assert "FORBIDDEN_" not in text
        # The full Q7 authored review text is retained separately. It may describe
        # scores; no scoring fields, policies or test expectations are exported.
        for key in ('"score":', '"formula":', '"classification":', '"expected_option_id":'):
            assert key not in text


def test_display_order_is_independent_of_physical_rows(real_source):
    def reverse_rows(workbook):
        for name in ["Questions", "Answer_Options"]:
            sheet = workbook[name]
            rows = list(sheet.values)[1:]
            for row_number, values in enumerate(reversed(rows), start=2):
                for column, value in enumerate(values, start=1):
                    sheet.cell(row_number, column).value = value
    original, _ = parse_workbook(real_source.read_bytes())
    reordered, _ = parse_workbook(edited_bytes(real_source, reverse_rows))
    assert original.catalog == reordered.catalog
    assert original.review == reordered.review


@pytest.mark.parametrize("status", ["Draft", "Unapproved", "Retired"])
def test_nonapproved_explanations_are_not_exported(real_source, status):
    raw = edited_bytes(real_source, lambda w: setattr(w["Explanations"]["F7"], "value", status))
    _, explanations = parse_workbook(raw)
    assert "EX06" not in {record.content_id for record in explanations.records}


def test_drafts_are_excluded_and_q7_never_retrieved(real_source, tmp_path):
    raw = edited_bytes(real_source, lambda w: setattr(w["Explanations"]["F7"], "value", "Draft"))
    catalog, explanations = parse_workbook(raw)
    assert "EX06" not in {r.content_id for r in explanations.records}
    retriever = CatalogRetriever(RELEASE)
    provider = JsonCatalogProvider(RELEASE)
    catalog = provider.load()
    result = retriever.retrieve(text="savings", question=catalog.questions[3], catalog=catalog)
    assert {r.sourceId for r in result} == {"EX06", "EX07"}
    assert all("overall score" not in r.text for r in result)
    with pytest.raises(ValueError):
        retriever.retrieve(text="confirm", question=provider.review, catalog=catalog)


@pytest.mark.parametrize("sheet,cell,value", [
    ("Questions", "A3", "Q1"), ("Questions", "B3", 1),
    ("Questions", "H2", "true"), ("Questions", "I3", "other-version"),
    ("Questions", "D2", None), ("Questions", "D2", " rewritten "),
    ("Questions", "D2", "=1+1"), ("Questions", "D1", "unknown_header"),
    ("Answer_Options", "A3", "Q1_A"), ("Answer_Options", "B2", "UNKNOWN"),
    ("Answer_Options", "C3", 1), ("Answer_Options", "E2", None),
    ("Explanations", "A3", "EX01"), ("Explanations", "B2", "UNKNOWN"),
    ("Explanations", "I2", "other-version"), ("Explanations", "F2", "Approvd"),
    ("Explanations", "D2", "x" * 601),
])
def test_invalid_workbook_does_not_publish(real_source, tmp_path, sheet, cell, value):
    raw = edited_bytes(real_source, lambda w: setattr(w[sheet][cell], "value", value))
    source = tmp_path / WORKBOOK_NAME
    source.write_bytes(raw)
    with pytest.raises(ValueError):
        import_workbook(source, output_root=tmp_path / "catalog", source_root=tmp_path / "source")
    assert not (tmp_path / "catalog").exists()
    assert not (tmp_path / "source").exists()


def test_approval_is_explicit_and_same_version_cannot_be_overwritten(real_source, tmp_path):
    catalog, _ = parse_workbook(real_source.read_bytes())
    assert catalog.approval.approved is False and catalog.catalog.financeApproved is False
    with pytest.raises(ValueError, match="basis"):
        parse_workbook(real_source.read_bytes(), finance_approved=True)
    target = import_workbook(real_source, output_root=tmp_path / "catalog", source_root=tmp_path / "source")
    before = (target / "catalog.json").read_bytes()
    with pytest.raises(ValueError, match="refusing overwrite"):
        import_workbook(real_source, output_root=tmp_path / "catalog", source_root=tmp_path / "source",
                        finance_approved=True, approval_basis="SYNTHETIC test approval, not a new real release")
    assert (target / "catalog.json").read_bytes() == before


def test_real_q4_mocked_adapter_and_service_keep_confirmation_boundary(repository):
    provider = JsonCatalogProvider(RELEASE)
    catalog = provider.load()
    observed = []
    def handler(request):
        data = json.loads(json.loads(request.content)["messages"][1]["content"])
        question = data["question"]
        if question["id"] == "Q4":
            assert question["options"] == [o.model_dump() for o in catalog.questions[3].options]
            assert {r["sourceId"] for r in data["context"]} == {"EX06", "EX07"}
            observed.append(question["id"])
        option = "Q4_D" if question["id"] == "Q4" else question["options"][0]["id"]
        return httpx.Response(200, json={"choices": [{"finish_reason": "stop", "message": {
            "role": "assistant", "content": json.dumps({"kind": "proposal", "optionId": option})}}]})
    model = OmlxAdapter(base_url="http://mock.test/v1", model="MOCK", api_key="MOCK",
                        transport=httpx.MockTransport(handler))
    with TestClient(create_app(repository=repository, catalog=provider,
                              retrieval=CatalogRetriever(RELEASE), model=model)) as client:
        journey = Journey(client)
        journey.send("start")
        for question in catalog.questions:
            response = journey.send("message", message={"role": "user", "text": "Mocked fictional answer"})
            assert all(a["questionId"] != question.id for a in response["confirmedAnswers"])
            proposal = response["proposal"]
            if question.id == "Q4":
                assert journey.post(journey.body("confirm", questionId="Q4", optionId="Q4_D")).status_code == 422
            journey.send("confirm", questionId=question.id, optionId=proposal["optionId"],
                         explicitConfirmation=question.requiresExplicitConfirmation)
            assert repository.list_for_session(journey.session_id) == []
        assert observed == ["Q4"]
        assert journey.send("save")["saveStatus"] == "saved"
        profile = repository.list_for_session(journey.session_id)[0]
        assert profile.score.status == "not_configured"
        assert [a.questionId for a in profile.answers] == [f"Q{i}" for i in range(1, 7)]
